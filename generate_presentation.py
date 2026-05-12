from __future__ import annotations

from pathlib import Path

from pptx import Presentation


def add_title_slide(prs: Presentation, title: str, subtitle: str) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = title
    slide.placeholders[1].text = subtitle


def add_bullets_slide(prs: Presentation, title: str, bullets: list[str]) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = title
    body = slide.shapes.placeholders[1].text_frame
    body.clear()
    for idx, line in enumerate(bullets):
        if idx == 0:
            p = body.paragraphs[0]
        else:
            p = body.add_paragraph()
        p.text = line


def main() -> None:
    prs = Presentation()

    add_title_slide(
        prs,
        "MoodWave: ML Conversational Music Recommender",
        "AMLL Project Presentation | April 2026",
    )

    add_bullets_slide(
        prs,
        "Slide 2: Problem Statement",
        [
            "Users want music that matches how they feel right now.",
            "Most apps require manual searching and mood selection.",
            "Goal: build a chatbot that understands emotion from text and recommends songs automatically.",
        ],
    )

    add_bullets_slide(
        prs,
        "Slide 3: Project Objective",
        [
            "Create a fully local ML-based conversational recommender.",
            "Detect tone from user text.",
            "Generate a context-aware chatbot response.",
            "Recommend songs and similar songs from a local dataset.",
            "Run offline without external APIs.",
        ],
    )

    add_bullets_slide(
        prs,
        "Slide 4: Final Architecture",
        [
            "Frontend: HTML + CSS + JavaScript",
            "Backend: Flask API",
            "ML Services: ToneService, CakeChatService, LastFMService",
            "Data: emotion CSV splits + local_songs.csv",
            "Artifacts: emotion_model.pkl, response_model.pkl, song_model.pkl",
        ],
    )

    add_bullets_slide(
        prs,
        "Slide 5: End-to-End Flow (/chat)",
        [
            "1) User message from UI",
            "2) ToneService predicts mood",
            "3) Tone mapped to chatbot emotion + song tag",
            "4) CakeChatService generates local reply",
            "5) LastFMService ranks songs",
            "6) API returns tone + reply + songs",
        ],
    )

    add_bullets_slide(
        prs,
        "Slide 6: ML Models Used",
        [
            "Emotion Model: TF-IDF + KMeans clustering + cluster-to-tone mapping",
            "Response Model: TF-IDF retrieval over curated prompt-response corpus",
            "Song Model: TF-IDF ranking and vector similarity over local songs",
            "Training Script: scripts/train_emotion_model.py",
        ],
    )

    add_bullets_slide(
        prs,
        "Slide 7: APIs Implemented",
        [
            "GET /",
            "POST /tone",
            "POST /response",
            "GET /songs?tag=...",
            "GET /similarsongs?track=...&artist=...",
            "POST /chat",
        ],
    )

    add_bullets_slide(
        prs,
        "Slide 8: Dataset and Artifacts",
        [
            "Emotion data: emotion_train.csv, emotion_val.csv, emotion_test.csv",
            "Song data: local_songs.csv",
            "Saved artifacts: model/emotion_model.pkl",
            "Saved artifacts: model/response_model.pkl",
            "Saved artifacts: model/song_model.pkl",
        ],
    )

    add_bullets_slide(
        prs,
        "Slide 9: Reliability and Fallback",
        [
            "Explicit phrase detection improves tone accuracy.",
            "If model artifact is missing, runtime training can initialize emotion model.",
            "If song model is unavailable, CSV fallback still returns recommendations.",
            "System remains usable in degraded environments.",
        ],
    )

    add_bullets_slide(
        prs,
        "Slide 10: Key Modules",
        [
            "app/routes.py: endpoint orchestration",
            "app/services/tone_service.py: tone inference",
            "app/services/cakechat_service.py: response generation",
            "app/services/lastfm_service.py: song ranking and similar songs",
            "scripts/train_emotion_model.py: training pipeline",
        ],
    )

    add_bullets_slide(
        prs,
        "Slide 11: Demo Script",
        [
            "Input: 'feeling very happy' -> tone: joy -> upbeat songs",
            "Input: 'i feel sad today' -> tone: sadness -> acoustic songs",
            "Input: 'i am angry right now' -> tone: anger -> rock songs",
            "Input: 'i feel neutral' -> tone: neutral -> chill songs",
        ],
    )

    add_bullets_slide(
        prs,
        "Slide 12: Tech Stack",
        [
            "Python 3.13",
            "Flask 3.1.1",
            "scikit-learn 1.5.2",
            "pandas 2.2.3",
            "numpy 2.1.3",
            "joblib 1.4.2",
        ],
    )

    add_bullets_slide(
        prs,
        "Slide 13: Impact and Use Cases",
        [
            "Mood-based personal music companion",
            "Conversational recommendation assistant",
            "Offline-capable prototype for emotion-aware UX",
            "Base architecture for healthcare/wellbeing adaptation",
        ],
    )

    add_bullets_slide(
        prs,
        "Slide 14: Conclusion",
        [
            "Project successfully transformed into fully ML-based system.",
            "All core functions run locally with stable fallbacks.",
            "API and UI are integrated and presentation-ready.",
            "Future scope: larger datasets, deep learning models, personalization.",
        ],
    )

    out_path = Path(__file__).resolve().parent.parent / "AMLL_Project_Presentation.pptx"
    prs.save(out_path)
    print(f"Saved: {out_path}")


if __name__ == "__main__":
    main()
